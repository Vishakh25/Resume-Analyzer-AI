# resume_analyzer_backend/analyzer/views.py
import os
import io
import json # Ensure json module is imported for parsing
import logging

# --- Django REST Framework Imports ---
from rest_framework.views import APIView # <--- THIS IS THE CRUCIAL LINE ADDED/CONFIRMED
from rest_framework.response import Response
from rest_framework import status
from rest_framework.parsers import MultiPartParser, FormParser

# --- Document Parsing Libraries ---
# Note: PyPDF2 is often preferred, but sometimes pypdf is also used.
# If you used PyPDF2, ensure it's `from PyPDF2 import PdfReader`
# If you installed pypdf, ensure it's `from pypdf import PdfReader`
from PyPDF2 import PdfReader # Assuming you installed PyPDF2
from docx import Document

# --- Google Gemini API ---
import google.generativeai as genai
from decouple import config # For environment variables

# Set up logging for better debugging in the console
logger = logging.getLogger(__name__)
# Optional: Configure logging if you want more verbose output
# logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')


# Configure the Google Gemini API with your API key from .env
# This block runs when the server starts
try:
    gemini_api_key = config('GOOGLE_API_KEY')
    if not gemini_api_key:
        raise ValueError("GOOGLE_API_KEY not found in .env or environment variables.")
    genai.configure(api_key=gemini_api_key)
    logger.info("Google Gemini API configured successfully.")
except Exception as e:
    logger.error(f"FATAL: Error configuring Google Gemini API. Server might not function correctly: {e}")
    # In production, you might want to stop the server here if API key is absolutely critical


class AnalyzeResumeView(APIView):
    """
    API View to handle resume upload and analysis using Google Gemini.
    Both post and _get_gemini_analysis are now synchronous.
    """
    parser_classes = (MultiPartParser, FormParser)

    # post method is now synchronous (no async def)
    def post(self, request, *args, **kwargs):
        logger.info("Received POST request for resume analysis.")

        # 1. Get the resume file from the request
        resume_file = request.FILES.get('resume_file')
        if not resume_file:
            logger.warning("No resume file provided in the request.")
            return Response(
                {"success": False, "message": "No resume file provided. Please upload a PDF or DOCX."},
                status=status.HTTP_400_BAD_REQUEST
            )

        # 2. Extract job description and other parameters from form data
        job_description = request.data.get('job_description', '')
        job_category = request.data.get('job_category', 'General')
        specific_role = request.data.get('specific_role', 'Applicant') # Default to 'Applicant' if not provided
        logger.info(f"Job Description length: {len(job_description)} chars.")
        logger.info(f"Job Category: {job_category}, Specific Role: {specific_role}")

        # 3. Extract plain text content from the uploaded resume file
        resume_text = self._extract_text_from_resume(resume_file)
        if not resume_text:
            logger.error("Failed to extract text from resume.")
            return Response(
                {"success": False, "message": "Could not extract text from resume file. Please ensure it's a valid PDF or DOCX and not scanned image."},
                status=status.HTTP_400_BAD_REQUEST
            )
        logger.info(f"Extracted {len(resume_text)} characters from resume.")

        # 4. Call the Gemini API to get the detailed analysis report
        try:
            # CHANGE: Directly call _get_gemini_analysis (it's now synchronous)
            report = self._get_gemini_analysis(resume_text, job_description, job_category, specific_role)
            logger.info("Successfully received analysis from Gemini API.")
            return Response(
                {"success": True, "message": "Analysis complete.", "report": report},
                status=status.HTTP_200_OK
            )
        except Exception as e:
            # Catch any exceptions during the Gemini API call or response parsing
            logger.exception("Error during Gemini API call:") # Logs the full traceback
            return Response(
                {"success": False, "message": f"An error occurred during analysis: {e}. Check server logs for details."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    def _extract_text_from_resume(self, resume_file):
        """
        Helper method to extract text content from PDF and DOCX files.
        """
        file_extension = os.path.splitext(resume_file.name)[1].lower()
        text = ""

        try:
            # Read the file content into a BytesIO object for PyPDF2/docx
            file_content_io = io.BytesIO(resume_file.read())

            if file_extension == '.pdf':
                reader = PdfReader(file_content_io)
                for page in reader.pages:
                    extracted_page_text = page.extract_text()
                    if extracted_page_text: # Ensure text was actually extracted from the page
                        text += extracted_page_text + "\n"
            elif file_extension == '.docx':
                document = Document(file_content_io)
                for paragraph in document.paragraphs:
                    text += paragraph.text + "\n"
            else:
                logger.warning(f"Unsupported file type for text extraction: {file_extension}")
                return None # Return None if file type is not supported
        except Exception as e:
            logger.error(f"Error extracting text from {resume_file.name} ({file_extension}): {e}")
            # Re-raise or return None based on desired error handling
            return None
        return text.strip() # Return stripped text to remove leading/trailing whitespace

    # CHANGE: _get_gemini_analysis method is now synchronous (no async def)
    def _get_gemini_analysis(self, resume_text, job_description, job_category, specific_role):
        """
        Sends resume and job description details to the Gemini API for comprehensive analysis.
        Uses the synchronous Gemini API call.
        """
        # Specify the Gemini model.
        # Ensure you use the synchronous `generate_content` method now.
        model = genai.GenerativeModel(model_name="gemini-2.0-flash")

        prompt = f"""
        You are an expert Resume Analyzer and ATS (Applicant Tracking System) specialist.
        Your task is to analyze the given resume against the provided job description (if available), job category, and specific job role.
        Provide a comprehensive and actionable analysis in a structured JSON format.

        **Input Data:**
        - **Job Category:** {job_category}
        - **Specific Job Role:** {specific_role}
        - **Job Description (if provided):**
          {job_description if job_description else "No specific job description provided. Analyze based on general industry standards for the specified Job Category and Specific Job Role."}
        - **Resume Content:**
          {resume_text}

        ---

        **Please provide the analysis in a structured JSON format with the following keys and detailed content. Ensure all fields are present, even if empty arrays or "N/A" text. Generate natural language text for string fields and lists for array fields. Focus heavily on actionable advice for ATS optimization and content alignment.**

        {{
            "job_role_analyzed": "{specific_role}",  # Dynamically includes the role provided as input
            "overall_fit_score": "A numerical score (0-100) indicating the overall suitability and direct match of the resume for the specified job role and description. Higher score means better alignment with key requirements.",
            "summary_of_fit": "A concise (2-4 sentences) general paragraph summarizing the resume's overall strengths and main areas for improvement related to the target role. Focus on how well it presents the candidate for the specific job.",

            "professional_profile_analysis": "A detailed analysis (2-3 paragraphs) of the resume's professional summary/objective and overall presentation. Evaluate its conciseness, impact, and immediate relevance to the target job. Advise on clarity, keyword integration, and whether it effectively hooks a recruiter.",
            "skills_analysis": "A detailed analysis (2-3 paragraphs) of the skills section. Assess the relevance, breadth, and depth of listed skills against the job role and description. Highlight any critical missing skills or areas where skill demonstration could be stronger (e.g., through projects or achievements). Suggest integrating skills within experience descriptions for better context.",
            "experience_analysis": "A detailed analysis (3-4 paragraphs) of the work experience section. Evaluate whether achievements are quantifiable and directly relevant to the job. Provide specific advice on transforming responsibilities into accomplishments using the STAR method (Situation, Task, Action, Result) to demonstrate impact. Comment on the use of action verbs and keywords within descriptions.",
            "education_analysis": "A detailed analysis (1-2 paragraphs) of the education section. Comment on relevance, completeness, and any notable academic achievements, relevant coursework, or certifications. Suggest how to align educational background better with the specific career goal if needed.",

            "ats_optimization_score": "A numerical score (0-100) indicating how well the resume is optimized for Applicant Tracking Systems. Consider keyword density, standard section headers, font choices (implied by parse-ability), use of bullet points vs. paragraphs, and general resume scannability. Higher score means easier parsing by ATS.",
            "ats_optimization_assessment": "A detailed paragraph (2-3 paragraphs) explaining the ATS optimization score. Highlight specific formatting issues that might hinder ATS parsing (e.g., complex layouts, heavy graphics, non-standard fonts that convert poorly) and suggest immediate, actionable fixes. Emphasize the importance of clear, parsable sections and relevant keyword usage for ATS.",

            "role_alignment_analysis": "A detailed analysis (2-3 paragraphs) specifically assessing how well the resume's content directly aligns with the duties, responsibilities, and qualifications outlined in the job description or standard expectations for the specified role. Provide concrete examples of strong alignment and significant misalignments, offering tailored advice on how to bridge those gaps effectively.",

            "strengths": [
                "Specific, strong relevant skills/experiences identified from the resume that directly match job requirements. (e.g., 'Proficiency in Python and machine learning, directly addressing AI Engineer requirements.')",
                "Demonstrated impact through quantifiable achievements using strong action verbs. (e.g., 'Increased data processing efficiency by 15% in previous role.')",
                "Clear and standard resume formatting that is ATS-friendly."
            ],
            "areas_for_improvement": [
                "Actionable advice on integrating more specific keywords from the job description into experience bullet points.",
                "Suggest adding a 'Projects' section to showcase practical application of skills.",
                "Advise on quantifying achievements with numbers, percentages, or results where currently vague.",
                "Recommend clarifying career objective/summary to be more role-specific.",
                "Suggest using more robust action verbs at the start of bullet points.",
                "Consider reordering sections to highlight most relevant experience first."
            ],
            "missing_keywords": [
                "Crucial keyword from job description (e.g., 'Containerization' if it's key and absent)",
                "Important skill from role expectations (e.g., 'Cloud Security Best Practices' if relevant and not found)"
            ],
            "suggested_summary": "A concise (2-3 sentences) professional summary suggestion tailored specifically for the provided job role, leveraging the resume's strengths and incorporating key terms for impact. This summary should act as a compelling elevator pitch.",
            "recommended_courses": [
                {{ "name": "Course Title 1 (e.g., 'Advanced React Development')", "provider": "Platform (e.g., Coursera, Udemy)", "link": "https://example.com/course1" }},
                {{ "name": "Course Title 2 (e.g., 'Certified Cloud Practitioner')", "provider": "Platform (e.g., AWS, edX)", "link": "https://example.com/course2" }}
            ],
            "ats_keywords_found": [
                "List of top 10-15 highly relevant keywords from the job description (or role expectations) found in the resume, indicating good ATS match. (e.g., 'JavaScript', 'SQL', 'Agile Methodologies')"
            ],
            "ats_keywords_missing": [
                "List of top 10-15 important keywords from the job description (or role expectations) *not* found in the resume, that are critical for ATS matching and human review. (e.g., 'CI/CD Pipelines', 'Microservices Architecture', 'AWS Lambda')"
            ],
            "potential_interview_questions": [
                "Question 1: Based on strengths/gaps, e.g., 'Can you elaborate on your experience with X, as mentioned in the job description, and how it directly applies to this role?'",
                "Question 2: 'Your resume mentions Y, but the role places a strong emphasis on Z. How do you plan to develop your skills in Z to meet the job requirements?'",
                "Question 3: 'Describe a situation where you faced a significant technical challenge and how you overcame it, demonstrating your problem-solving skills and resilience.'",
                "Question 4: 'Based on your experience, how do you ensure your projects are completed on time and within budget?'"
            ]
        }}

        Ensure the response is valid JSON and directly follows the specified structure. Do not include any additional text or formatting outside the JSON object.
        If any section's data is not applicable or cannot be generated, use appropriate placeholders like 'N/A', 0, or empty arrays to maintain the JSON structure.
        The overall_fit_score and ats_optimization_score should always be numerical values between 0 and 100.
        """
        logger.info("Sending prompt to Gemini API with detailed request...")
        response = model.generate_content(prompt) # Synchronous call
        text_response = response.text

        if text_response.strip().startswith("```json"):
            json_start = text_response.find("```json") + len("```json")
            json_end = text_response.rfind("```")
            if json_end > json_start:
                text_response = text_response[json_start:json_end].strip()
            else:
                logger.warning("JSON markdown block started but not properly closed. Attempting to parse full text.")
                text_response = text_response.replace("```json", "").replace("```", "").strip()
        elif text_response.strip().startswith("```"):
            json_start = text_response.find("```") + len("```")
            json_end = text_response.rfind("```")
            if json_end > json_start:
                text_response = text_response[json_start:json_end].strip()
            else:
                logger.warning("Generic markdown block started but not properly closed. Attempting to parse full text.")
                text_response = text_response.replace("```", "").strip()

        try:
            parsed_report = json.loads(text_response)
            return parsed_report
        except json.JSONDecodeError as e:
            logger.error(f"Failed to parse Gemini response as JSON: {e}")
            logger.error(f"Raw Gemini response (attempted parse):\n{text_response}")
            raise Exception(f"AI response format error: Could not parse AI response as JSON. Details: {e}. Raw response start: {text_response[:500]}...")